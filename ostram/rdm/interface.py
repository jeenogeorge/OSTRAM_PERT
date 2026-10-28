"""The RDM Excel interface: the template and the reader.

The workbook has six sheets:

``README``
    How to fill the workbook.
``Setup``
    One value per setting: the run name, the root scenario every future is
    built from, the number of samples per future, the random seed, whether
    each named future also runs without its sampled uncertainties, whether
    the shared VRE ceilings apply, and the model's last year.
``Futures``
    One row per named future: ``Future_ID``, ``Name``, ``Description`` and
    ``Run`` (Y/N).
``Future_Changes``
    One row per deterministic change of a named future.
``Uncertainty_Table``
    One row per ranged uncertainty, sampled between ``Min`` and ``Max`` and
    layered inside the futures named in ``Applies_To``.
``Outputs``
    One row per input parameter or result variable, with ``Export`` (Y/N)
    deciding whether it reaches the RDM CSVs.

A hidden ``Lists`` sheet holds the dropdown values.  :func:`read_interface`
validates every sheet and reports all problems at once, each with its sheet
and row, so a partner can fix the workbook in one pass.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FUTURE_ID = re.compile(r"^[A-Za-z][A-Za-z0-9]*$")
RUN_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]*$")

TECH_SHEETS = ("Secondary Techs", "Primary Techs", "Demand Techs", "VariableCost")
DEMAND_SHEETS = ("Demand_Projection",)
SHEETS = TECH_SHEETS + DEMAND_SHEETS
TECH_SELECTORS = ("Tech", "Techs", "Tech_Prefix")
FUEL_SELECTORS = ("Fuel", "Fuels", "Fuel_Prefix")
CHANGE_OPERATIONS = ("Multiply", "Multiply_Ramp", "Set", "Copy_From_Tech")
UNCERTAINTY_METHODS = ("Multiply", "Multiply_Ramp", "Set")
YES_NO = ("Y", "N")
CEILING_CHOICES = ("Auto", "Y", "N")

DEFAULT_ROOTS = ("A_Calibrated_BAU", "B_Optimised_VRE", "C_Target_VRE", "BAU")

# Parameters offered in the Parameter dropdowns, by sheet family.
TECH_PARAMETERS = (
    "CapitalCost", "FixedCost", "VariableCost", "ResidualCapacity",
    "TotalAnnualMaxCapacity", "TotalAnnualMinCapacity",
    "TotalAnnualMaxCapacityInvestment", "TotalAnnualMinCapacityInvestment",
    "TotalTechnologyAnnualActivityUpperLimit", "TotalTechnologyAnnualActivityLowerLimit",
    "AvailabilityFactor",
)
DEMAND_PARAMETERS = ("SpecifiedAnnualDemand",)

# The Outputs sheet lists every input parameter ("param") and result variable
# ("result") of the model's otoole conversion file, with the sets each is
# indexed by.  These are exported by default; every other variable is N, so
# the RDM CSVs stay as small as before unless the user ticks more.
DEFAULT_INPUT_EXPORTS = frozenset({
    "AvailabilityFactor", "CapitalCost", "CapitalCostStorage", "EmissionActivityRatio",
    "FixedCost", "OperationalLife", "ReserveMargin", "ResidualCapacity",
    "ResidualStorageCapacity", "SpecifiedAnnualDemand", "TotalAnnualMaxCapacity",
    "TotalAnnualMaxCapacityInvestment", "TotalAnnualMinCapacityInvestment",
    "TotalTechnologyAnnualActivityUpperLimit", "VariableCost",
})
DEFAULT_OUTPUT_EXPORTS = frozenset({
    "AccumulatedNewCapacity", "AnnualEmissions", "AnnualFixedOperatingCost",
    "AnnualTechnologyEmission", "AnnualVariableOperatingCost", "CapitalInvestment",
    "CapitalInvestmentStorage", "DiscountedCapitalInvestment", "DiscountedCostByTechnology",
    "DiscountedOperationalCost", "NewCapacity", "NewStorageCapacity", "OperatingCost",
    "ProductionByTechnologyAnnual", "TotalCapacityAnnual", "TotalDiscountedCost",
    "TotalTechnologyAnnualActivity",
})
CONVERSION_FORMAT = Path(__file__).resolve().parents[1] / "resources" / "compilation" / "conversion_format.yaml"


def _conversion_format() -> Path:
    """The active setup's otoole conversion file, else the package's own copy."""

    try:
        from ostram.paths import resolve_paths
        path = resolve_paths().compilation_resources / "conversion_format.yaml"
        if path.is_file():
            return path
    except Exception:  # no project or setup context: fall back to the package file
        pass
    return CONVERSION_FORMAT


def variable_catalogue() -> dict[str, dict[str, tuple[str, ...]]]:
    """Return ``{"Input": {name: indices}, "Output": {name: indices}}``, sorted by name.

    Read from the otoole conversion file, so the Outputs sheet offers exactly
    the parameters and results B2 can write to the per-future CSVs.
    """

    import yaml

    raw = yaml.safe_load(_conversion_format().read_text(encoding="utf-8")) or {}
    out: dict[str, dict[str, tuple[str, ...]]] = {"Input": {}, "Output": {}}
    for name, spec in sorted(raw.items(), key=lambda item: item[0].lower()):
        kind = {"param": "Input", "result": "Output"}.get((spec or {}).get("type"))
        if kind:
            out[kind][name] = tuple(spec.get("indices") or ())
    return out


def _defaults(kind: str) -> dict[str, str]:
    chosen = DEFAULT_INPUT_EXPORTS if kind == "Input" else DEFAULT_OUTPUT_EXPORTS
    return {name: ("Y" if name in chosen else "N") for name in variable_catalogue()[kind]}


# Name -> default export choice, for every variable of the model.
INPUT_VARIABLES = _defaults("Input")
OUTPUT_VARIABLES = _defaults("Output")

SETUP_ROWS = (
    ("Run_Name", "rdm_run_01", "Folder name of this run under <workspace>/rdm/. Letters, digits, _ and -."),
    ("Base_Scenario", "B_Optimised_VRE", "Root scenario every future is built from."),
    ("Samples_Per_Future", 0, "Samples of the ranged uncertainties inside each named future. 0 = no sampling."),
    ("Random_Seed", 12345, "Seed of the Latin Hypercube sample. Same seed, same futures."),
    ("Run_Central_Future", "Y", "Y also runs each named future without its sampled uncertainties (e.g. F01 beside F01_s001...)."),
    ("Apply_VRE_Ceilings", "Auto", "Shared VRE ceilings on every future. Auto = Y when the setup defines them."),
    ("Model_First_Year", 2023, "First model year. A change with only Year_To starts here."),
    ("Model_Last_Year", 2050, "Last model year. A change with only Year_From runs to here; Multiply_Ramp holds its final factor up to here."),
)

CHANGE_COLUMNS = (
    "Future_ID", "Sheet", "Parameter", "Selector_Type", "Selector",
    "Operation", "Value", "Year_From", "Year_To", "Note",
)
UNCERTAINTY_COLUMNS = (
    "X_ID", "Name", "Applies_To", "Sheet", "Parameter", "Selector_Type", "Selector",
    "Method", "Min", "Max", "Year_From", "Year_To", "Linked_To", "Note",
)
FUTURE_COLUMNS = ("Future_ID", "Name", "Description", "Run")
OUTPUT_COLUMNS = ("Kind", "Variable", "Export")
# Written for the reader's benefit only; read_interface ignores it.
OUTPUT_INFO_COLUMN = "Indexed_By"


class InterfaceError(ValueError):
    """The interface workbook has one or more problems; the message lists them all."""

    def __init__(self, problems: Sequence[str]):
        self.problems = list(problems)
        super().__init__(
            "The RDM interface has "
            f"{len(self.problems)} problem{'s' if len(self.problems) != 1 else ''}:\n"
            + "\n".join(f"  - {p}" for p in self.problems)
        )


@dataclass(frozen=True)
class Change:
    """One deterministic change, or one uncertainty with its sampled value filled in."""

    sheet: str
    parameter: str
    selector_type: str
    selector: str
    operation: str
    value: object
    year_from: int | None = None
    year_to: int | None = None
    note: str = ""
    where: str = ""


@dataclass(frozen=True)
class Uncertainty:
    """One ranged uncertainty of the Uncertainty_Table sheet."""

    x_id: str
    name: str
    applies_to: tuple[str, ...] | None  # None means every future
    sheet: str
    parameter: str
    selector_type: str
    selector: str
    method: str
    minimum: float
    maximum: float
    year_from: int | None
    year_to: int | None
    linked_to: str | None
    note: str
    where: str

    def change(self, value: float) -> Change:
        """Return this uncertainty as a change at the sampled ``value``."""

        return Change(
            sheet=self.sheet, parameter=self.parameter,
            selector_type=self.selector_type, selector=self.selector,
            operation=self.method, value=value,
            year_from=self.year_from, year_to=self.year_to,
            note=f"{self.x_id} {self.name}".strip(), where=self.where,
        )


@dataclass(frozen=True)
class Future:
    """One named future of the Futures sheet."""

    future_id: str
    name: str
    description: str
    run: bool


@dataclass
class Interface:
    """A validated interface workbook."""

    path: Path
    run_name: str
    base_scenario: str
    samples_per_future: int
    random_seed: int
    run_central_future: bool
    apply_vre_ceilings: str
    model_first_year: int
    model_last_year: int
    futures: list[Future]
    changes: dict[str, list[Change]]
    uncertainties: list[Uncertainty]
    input_variables: list[str] = field(default_factory=list)
    output_variables: list[str] = field(default_factory=list)

    @property
    def active_futures(self) -> list[Future]:
        """Return the futures whose Run column is Y, in sheet order."""

        return [future for future in self.futures if future.run]


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_NOTE_FONT = Font(italic=True, color="595959")


def _header(ws, columns: Sequence[str], widths: Sequence[int]) -> None:
    for index, (name, width) in enumerate(zip(columns, widths), start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def _list_column(ws, column: int, title: str, values: Iterable[object]) -> str:
    ws.cell(row=1, column=column, value=title)
    values = list(values)
    for row, value in enumerate(values, start=2):
        ws.cell(row=row, column=column, value=value)
    letter = get_column_letter(column)
    return f"Lists!${letter}$2:${letter}${max(len(values) + 1, 2)}"


def _validation(ws, formula: str, cells: str, *, strict: bool = True) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if not strict:
        # Allow free text too (several techs, a source tech, an ALL keyword...).
        dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add(cells)


def write_template(
    path: Path | str,
    *,
    roots: Sequence[str] = DEFAULT_ROOTS,
    base_scenario: str | None = None,
    technologies: Sequence[str] = (),
    fuels: Sequence[str] = (),
    model_first_year: int = 2023,
    model_last_year: int = 2050,
    examples: bool = True,
) -> Path:
    """Write a blank interface workbook to ``path`` and return the path.

    ``technologies`` and ``fuels`` fill the selector dropdowns when they are
    known (``ostram rdm template`` reads them from the prepared root).  With
    ``examples`` the sheets carry a worked example that a user overwrites.
    """

    path = Path(path)
    wb = Workbook()

    readme = wb.active
    readme.title = "README"
    readme.column_dimensions["A"].width = 110
    for row, line in enumerate(README_LINES, start=1):
        cell = readme.cell(row=row, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row == 1:
            cell.font = Font(bold=True, size=14)
        elif line and line == line.upper() and not line.startswith(" "):
            cell.font = Font(bold=True)

    lists = wb.create_sheet("Lists")
    roots_ref = _list_column(lists, 1, "Roots", roots)
    sheets_ref = _list_column(lists, 2, "Sheets", SHEETS)
    params_ref = _list_column(lists, 3, "Parameters", TECH_PARAMETERS + DEMAND_PARAMETERS)
    selectors_ref = _list_column(lists, 4, "Selector_Types", TECH_SELECTORS + FUEL_SELECTORS)
    operations_ref = _list_column(lists, 5, "Operations", CHANGE_OPERATIONS)
    methods_ref = _list_column(lists, 6, "Methods", UNCERTAINTY_METHODS)
    yes_no_ref = _list_column(lists, 7, "Y_N", YES_NO)
    ceilings_ref = _list_column(lists, 8, "Ceilings", CEILING_CHOICES)
    selector_ref = _list_column(lists, 9, "Technologies_and_Fuels", [*technologies, *fuels])
    lists.sheet_state = "hidden"

    setup = wb.create_sheet("Setup")
    _header(setup, ("Setting", "Value", "Meaning"), (24, 22, 100))
    defaults = {
        "Base_Scenario": base_scenario or ("B_Optimised_VRE" if "B_Optimised_VRE" in roots else roots[0]),
        "Model_First_Year": model_first_year,
        "Model_Last_Year": model_last_year,
    }
    for row, (key, default, meaning) in enumerate(SETUP_ROWS, start=2):
        setup.cell(row=row, column=1, value=key).font = Font(bold=True)
        setup.cell(row=row, column=2, value=defaults.get(key, default))
        setup.cell(row=row, column=3, value=meaning).font = _NOTE_FONT
    _validation(setup, roots_ref, "B3")
    _validation(setup, yes_no_ref, "B6")
    _validation(setup, ceilings_ref, "B7")

    futures = wb.create_sheet("Futures")
    _header(futures, FUTURE_COLUMNS, (12, 28, 70, 8))
    _validation(futures, yes_no_ref, "D2:D500")

    changes = wb.create_sheet("Future_Changes")
    _header(changes, CHANGE_COLUMNS, (12, 18, 38, 14, 26, 16, 12, 11, 11, 50))
    _validation(changes, sheets_ref, "B2:B2000")
    _validation(changes, params_ref, "C2:C2000")
    _validation(changes, selectors_ref, "D2:D2000")
    _validation(changes, selector_ref, "E2:E2000", strict=False)
    _validation(changes, operations_ref, "F2:F2000")

    uncertainty = wb.create_sheet("Uncertainty_Table")
    _header(uncertainty, UNCERTAINTY_COLUMNS, (8, 26, 14, 18, 38, 14, 26, 16, 9, 9, 11, 11, 11, 50))
    _validation(uncertainty, sheets_ref, "D2:D2000")
    _validation(uncertainty, params_ref, "E2:E2000")
    _validation(uncertainty, selectors_ref, "F2:F2000")
    _validation(uncertainty, selector_ref, "G2:G2000", strict=False)
    _validation(uncertainty, methods_ref, "H2:H2000")

    outputs = wb.create_sheet("Outputs")
    _header(outputs, (*OUTPUT_COLUMNS, OUTPUT_INFO_COLUMN), (10, 44, 9, 58))
    catalogue = variable_catalogue()
    row = 2
    for kind, variables in (("Input", INPUT_VARIABLES), ("Output", OUTPUT_VARIABLES)):
        for name, export in variables.items():
            outputs.cell(row=row, column=1, value=kind)
            outputs.cell(row=row, column=2, value=name)
            outputs.cell(row=row, column=3, value=export)
            outputs.cell(row=row, column=4, value=", ".join(catalogue[kind].get(name, ())))
            row += 1
    outputs.auto_filter.ref = f"A1:D{row - 1}"
    _validation(outputs, yes_no_ref, f"C2:C{row}")

    if examples:
        _write_examples(futures, changes, uncertainty, technologies, fuels)

    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _pick(names: Sequence[str], prefix: str, fallback: str) -> str:
    return next((n for n in names if n.startswith(prefix)), fallback)


def _write_examples(futures, changes, uncertainty, technologies, fuels) -> None:
    solar = "PWRSPV"
    coal_fuel_tech = _pick(technologies, "MINCOA", "MINCOABGD")
    rows = [
        ("F00", "Baseline", "The root scenario with no changes.", "Y"),
        ("F01", "Cheap solar", "Solar capital costs fall 30% by 2035.", "Y"),
        ("F02", "High demand, dear coal", "Demand 15% higher; coal price doubled.", "Y"),
    ]
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            futures.cell(row=r, column=c, value=value)
    change_rows = [
        ("F01", "Secondary Techs", "CapitalCost", "Tech_Prefix", solar, "Multiply_Ramp", 0.7, 2026, 2035,
         "ASSUMPTION: example, solar cost falls to 70% by 2035"),
        ("F02", "Demand_Projection", "SpecifiedAnnualDemand", "Fuel_Prefix", "ELC", "Multiply", 1.15, 2027, None,
         "ASSUMPTION: example, demand 15% above the root from 2027"),
        ("F02", "VariableCost", "VariableCost", "Tech", coal_fuel_tech, "Multiply", 2.0, None, None,
         "ASSUMPTION: example, coal price doubled"),
    ]
    for r, values in enumerate(change_rows, start=2):
        for c, value in enumerate(values, start=1):
            changes.cell(row=r, column=c, value=value)
    uncertainty_rows = [
        ("X1", "Demand growth", "ALL", "Demand_Projection", "SpecifiedAnnualDemand", "Fuel_Prefix", "ELC",
         "Multiply_Ramp", 0.9, 1.1, 2026, 2050, None, "Demand from 90% to 110% of the future's path by 2050"),
        ("X2", "Gas price", "ALL", "VariableCost", "VariableCost", "Tech_Prefix", "MINGAS",
         "Multiply", 0.8, 1.5, None, None, None, "Gas price from 80% to 150%"),
    ]
    for r, values in enumerate(uncertainty_rows, start=2):
        for c, value in enumerate(values, start=1):
            uncertainty.cell(row=r, column=c, value=value)


README_LINES = (
    "OSTRAM RDM interface",
    "",
    "Fill this workbook, then run:  python -m ostram [--setup <setup>] rdm run --interface <this file>",
    "Nothing in the model inputs or config/ is changed. Every future is built in <workspace>/rdm/<Run_Name>/.",
    "",
    "HOW A FUTURE IS BUILT",
    "  1. Start from the root scenario named in Setup > Base_Scenario (after the shared VRE ceilings, when they apply).",
    "  2. Apply the future's rows of Future_Changes, top to bottom.",
    "  3. When Samples_Per_Future > 0: apply the sampled uncertainties of Uncertainty_Table that apply to the future.",
    "     Each named future F01 then becomes F01_s001, F01_s002, ... (layered sampling).",
    "     Sample s001 uses the same uncertainty draw in every future, so futures are compared like for like.",
    "     With Run_Central_Future = Y, F01 itself (no sampled uncertainties) also runs.",
    "",
    "SETUP",
    "  One value per setting. See the Meaning column.",
    "",
    "FUTURES",
    "  Future_ID: a short code, a letter followed by letters or digits (F00, F01, HiDem). It becomes the scenario name.",
    "  Run: Y to run the future, N to keep it in the workbook without running it.",
    "  A future with no rows in Future_Changes is the root itself (a baseline).",
    "",
    "FUTURE_CHANGES  (one row per change)",
    "  Sheet: the A-O workbook sheet that holds the parameter: Secondary Techs (power plants, storage, interconnectors),",
    "         Primary Techs (fuel supply), Demand Techs (grid tiers), VariableCost (fuel prices), Demand_Projection (demand).",
    "  Parameter: the OSeMOSYS parameter, e.g. CapitalCost, TotalAnnualMaxCapacity, SpecifiedAnnualDemand.",
    "  Selector_Type / Selector:",
    "    Tech        one technology, e.g. PWRSPVBGDXX",
    "    Techs       several technologies separated by commas",
    "    Tech_Prefix every technology whose name starts with the text, e.g. PWRSPV (all solar)",
    "    Fuel / Fuels / Fuel_Prefix  the same for demand fuels on Demand_Projection, e.g. ELCBGDXX03",
    "  Operation and Value:",
    "    Multiply        multiply by Value. Year_From / Year_To limit the years; blank = every year.",
    "    Multiply_Ramp   multiplier rises in a straight line from 1 at Year_From to Value at Year_To, then stays at Value",
    "                    up to Model_Last_Year. Years before Year_From are unchanged. Both years are required.",
    "    Set             set to Value. Year_From / Year_To limit the years; blank = every year.",
    "    Copy_From_Tech  copy every year from the technology named in Value (same sheet and parameter).",
    "  Note: the basis of the change: a source (URL, DOI) or ASSUMPTION and a reason.",
    "",
    "UNCERTAINTY_TABLE  (one row per ranged uncertainty)",
    "  X_ID: a short code (X1, X2...). Name: a label.",
    "  Applies_To: ALL, or the Future_IDs it applies to, separated by commas.",
    "  Sheet, Parameter, Selector_Type, Selector, Year_From, Year_To: as in Future_Changes.",
    "  Method: Multiply, Multiply_Ramp or Set. The sampled value replaces Value.",
    "  Min / Max: the range. Values are drawn by Latin Hypercube sampling, uniform between Min and Max.",
    "  Linked_To: the X_ID of another row. The two rows then move together: the same position in their ranges.",
    "  Sampled uncertainties apply after the future's own changes, so a Multiply scales the future's path.",
    "",
    "OUTPUTS",
    "  Every input parameter and result variable of the model is listed, with the sets it is indexed by (Indexed_By).",
    "  Export Y puts it in the RDM CSVs. Filter the sheet on Indexed_By to find, e.g., the FUEL- or EMISSION-keyed ones.",
    "  Timeslice variables (TIMESLICE in Indexed_By) are large; they are N by default.",
    "  A ticked variable the model leaves empty (no value in any future) is reported by collect and left out of the CSV.",
    "",
    "WHAT THE RUN WRITES  (<workspace>/rdm/<Run_Name>/)",
    "  RDM_Experiment_Design.csv  one row per future: its named future, sample number and the sampled value of every uncertainty.",
    "  RDM_Inputs.csv             the selected input parameters of every future.",
    "  RDM_Outputs.csv            the selected result variables of every future.",
    "  RDM_Run_Status.csv         which futures solved.",
    "  scenarios/<future>/scenario.yaml  the edits each future applied, readable and re-runnable.",
)


# ---------------------------------------------------------------------------
# Reader
# ---------------------------------------------------------------------------
def _closest(name: str, choices: Iterable[str]) -> str | None:
    """The choice that ``name`` most likely misspells, ignoring case, or None."""

    import difflib

    choices = list(choices)
    by_lower = {c.lower(): c for c in choices}
    if name.lower() in by_lower:
        return by_lower[name.lower()]
    hits = difflib.get_close_matches(name.lower(), list(by_lower), n=1, cutoff=0.75)
    return by_lower[hits[0]] if hits else None


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _rows(ws, columns: Sequence[str]):
    """Yield ``(row_number, {column: value})`` for each non-blank row of ``ws``."""

    header = [_text(c.value) for c in ws[1]]
    index = {name: header.index(name) for name in columns if name in header}
    missing = [name for name in columns if name not in index]
    if missing:
        raise InterfaceError([f"{ws.title}: missing column(s) {', '.join(missing)}"])
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {name: row[i] if i < len(row) else None for name, i in index.items()}
        if all(_text(v) == "" for v in values.values()):
            continue
        yield number, values


def _int(value, where: str, name: str, problems: list[str], *, required=False) -> int | None:
    if _text(value) == "":
        if required:
            problems.append(f"{where}: {name} is required")
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        problems.append(f"{where}: {name} must be a whole number, not {value!r}")
        return None
    if number != int(number):
        problems.append(f"{where}: {name} must be a whole number, not {value!r}")
        return None
    return int(number)


def _float(value, where: str, name: str, problems: list[str]) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        problems.append(f"{where}: {name} must be a number, not {value!r}")
        return None
    if number != number or number in (float("inf"), float("-inf")):
        problems.append(f"{where}: {name} must be a finite number")
        return None
    return number


def _yes(value, where: str, name: str, problems: list[str], default: bool) -> bool:
    text = _text(value).upper()
    if text == "":
        return default
    if text in ("Y", "YES", "TRUE", "1"):
        return True
    if text in ("N", "NO", "FALSE", "0"):
        return False
    problems.append(f"{where}: {name} must be Y or N, not {value!r}")
    return default


def _check_target(where, sheet, parameter, selector_type, selector, problems) -> None:
    if sheet not in SHEETS:
        problems.append(f"{where}: Sheet {sheet!r} is not one of {', '.join(SHEETS)}")
        return
    if parameter == "":
        problems.append(f"{where}: Parameter is required")
    if sheet in DEMAND_SHEETS:
        if selector_type not in FUEL_SELECTORS:
            problems.append(f"{where}: on {sheet} the Selector_Type is Fuel, Fuels or Fuel_Prefix")
        if parameter and parameter not in DEMAND_PARAMETERS:
            problems.append(f"{where}: {sheet} holds only SpecifiedAnnualDemand")
    elif selector_type not in TECH_SELECTORS:
        problems.append(f"{where}: on {sheet} the Selector_Type is Tech, Techs or Tech_Prefix")
    if selector == "":
        problems.append(f"{where}: Selector is required")


def _check_years(where, operation, year_from, year_to, first_year, last_year, problems) -> None:
    if operation == "Multiply_Ramp":
        if year_from is None or year_to is None:
            problems.append(f"{where}: Multiply_Ramp needs both Year_From and Year_To")
            return
        if year_to <= year_from:
            problems.append(f"{where}: Year_To must be after Year_From")
    if year_from is not None and year_to is not None and year_to < year_from:
        problems.append(f"{where}: Year_To must not be before Year_From")
    for name, year in (("Year_From", year_from), ("Year_To", year_to)):
        if year is not None and year > last_year:
            problems.append(f"{where}: {name} {year} is after Model_Last_Year {last_year}")
        if year is not None and year < first_year:
            problems.append(f"{where}: {name} {year} is before Model_First_Year {first_year}")


def read_interface(path: Path | str) -> Interface:
    """Read and validate the interface workbook at ``path``.

    Raises :class:`InterfaceError` listing every problem found.
    """

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"interface workbook not found: {path}")
    wb = load_workbook(path, data_only=True, read_only=True)
    problems: list[str] = []
    required = ("Setup", "Futures", "Future_Changes", "Uncertainty_Table", "Outputs")
    missing = [name for name in required if name not in wb.sheetnames]
    if missing:
        raise InterfaceError([f"missing sheet(s): {', '.join(missing)}"])

    # Setup
    settings = {}
    for row in wb["Setup"].iter_rows(min_row=2, values_only=True):
        if row and _text(row[0]):
            settings[_text(row[0])] = row[1] if len(row) > 1 else None
    run_name = _text(settings.get("Run_Name"))
    if not RUN_NAME.fullmatch(run_name):
        problems.append(f"Setup: Run_Name {run_name!r} must use letters, digits, _ and - only")
    base = _text(settings.get("Base_Scenario"))
    if not base:
        problems.append("Setup: Base_Scenario is required")
    samples = _int(settings.get("Samples_Per_Future"), "Setup", "Samples_Per_Future", problems) or 0
    if samples < 0 or samples > 9999:
        problems.append("Setup: Samples_Per_Future must be between 0 and 9999")
    seed = _int(settings.get("Random_Seed"), "Setup", "Random_Seed", problems)
    central = _yes(settings.get("Run_Central_Future"), "Setup", "Run_Central_Future", problems, True)
    ceilings = (_text(settings.get("Apply_VRE_Ceilings")) or "Auto").upper()
    ceilings = {"AUTO": "Auto", "Y": "Y", "YES": "Y", "N": "N", "NO": "N"}.get(ceilings, ceilings)
    if ceilings not in CEILING_CHOICES:
        problems.append("Setup: Apply_VRE_Ceilings must be Auto, Y or N")
    first_year = _int(settings.get("Model_First_Year"), "Setup", "Model_First_Year", problems) or 2023
    last_year = _int(settings.get("Model_Last_Year"), "Setup", "Model_Last_Year", problems) or 2050
    if last_year <= first_year:
        problems.append("Setup: Model_Last_Year must be after Model_First_Year")

    # Futures
    futures: list[Future] = []
    seen: set[str] = set()
    try:
        for number, row in _rows(wb["Futures"], FUTURE_COLUMNS):
            where = f"Futures row {number}"
            fid = _text(row["Future_ID"])
            if not FUTURE_ID.fullmatch(fid):
                problems.append(f"{where}: Future_ID {fid!r} must be a letter followed by letters or digits")
                continue
            if fid in seen:
                problems.append(f"{where}: Future_ID {fid} is repeated")
                continue
            seen.add(fid)
            futures.append(Future(fid, _text(row["Name"]), _text(row["Description"]),
                                  _yes(row["Run"], where, "Run", problems, True)))
    except InterfaceError as error:
        problems.extend(error.problems)
    if futures and not any(f.run for f in futures):
        problems.append("Futures: no future has Run = Y")
    elif not futures:
        problems.append("Futures: list at least one future")

    # Future_Changes
    changes: dict[str, list[Change]] = {f.future_id: [] for f in futures}
    try:
        for number, row in _rows(wb["Future_Changes"], CHANGE_COLUMNS):
            where = f"Future_Changes row {number}"
            fid = _text(row["Future_ID"])
            if fid not in changes:
                # Keep checking the row, so all its problems surface in one pass.
                problems.append(f"{where}: Future_ID {fid!r} is not on the Futures sheet")
            sheet, parameter = _text(row["Sheet"]), _text(row["Parameter"])
            selector_type, selector = _text(row["Selector_Type"]), _text(row["Selector"])
            operation = _text(row["Operation"])
            _check_target(where, sheet, parameter, selector_type, selector, problems)
            if operation not in CHANGE_OPERATIONS:
                problems.append(f"{where}: Operation {operation!r} is not one of {', '.join(CHANGE_OPERATIONS)}")
                continue
            year_from = _int(row["Year_From"], where, "Year_From", problems)
            year_to = _int(row["Year_To"], where, "Year_To", problems)
            _check_years(where, operation, year_from, year_to, first_year, last_year, problems)
            if operation == "Copy_From_Tech":
                value = _text(row["Value"])
                if not value:
                    problems.append(f"{where}: Copy_From_Tech needs the source technology in Value")
                if sheet in DEMAND_SHEETS:
                    problems.append(f"{where}: Copy_From_Tech does not apply to {sheet}")
                if year_from is not None or year_to is not None:
                    problems.append(f"{where}: Copy_From_Tech copies every year; leave the years blank")
            else:
                value = _float(row["Value"], where, "Value", problems)
            if fid in changes:
                changes[fid].append(Change(sheet, parameter, selector_type, selector, operation, value,
                                           year_from, year_to, _text(row["Note"]), where))
    except InterfaceError as error:
        problems.extend(error.problems)

    # Uncertainty_Table
    uncertainties: list[Uncertainty] = []
    try:
        for number, row in _rows(wb["Uncertainty_Table"], UNCERTAINTY_COLUMNS):
            where = f"Uncertainty_Table row {number}"
            x_id = _text(row["X_ID"])
            if not FUTURE_ID.fullmatch(x_id):
                problems.append(f"{where}: X_ID {x_id!r} must be a letter followed by letters or digits")
                continue
            if any(u.x_id == x_id for u in uncertainties):
                problems.append(f"{where}: X_ID {x_id} is repeated")
                continue
            applies = _text(row["Applies_To"]) or "ALL"
            applies_to = None if applies.upper() == "ALL" else tuple(
                part.strip() for part in applies.split(",") if part.strip()
            )
            for fid in applies_to or ():
                if fid not in changes:
                    problems.append(f"{where}: Applies_To names {fid!r}, which is not on the Futures sheet")
            sheet, parameter = _text(row["Sheet"]), _text(row["Parameter"])
            selector_type, selector = _text(row["Selector_Type"]), _text(row["Selector"])
            _check_target(where, sheet, parameter, selector_type, selector, problems)
            method = _text(row["Method"])
            if method not in UNCERTAINTY_METHODS:
                problems.append(f"{where}: Method {method!r} is not one of {', '.join(UNCERTAINTY_METHODS)}")
            low = _float(row["Min"], where, "Min", problems)
            high = _float(row["Max"], where, "Max", problems)
            if low is not None and high is not None and high < low:
                problems.append(f"{where}: Max must not be below Min")
            year_from = _int(row["Year_From"], where, "Year_From", problems)
            year_to = _int(row["Year_To"], where, "Year_To", problems)
            _check_years(where, method, year_from, year_to, first_year, last_year, problems)
            linked = _text(row["Linked_To"]) or None
            uncertainties.append(Uncertainty(
                x_id, _text(row["Name"]), applies_to, sheet, parameter, selector_type, selector,
                method, low if low is not None else 0.0, high if high is not None else 0.0,
                year_from, year_to, linked, _text(row["Note"]), where,
            ))
    except InterfaceError as error:
        problems.extend(error.problems)
    ids = {u.x_id: u for u in uncertainties}
    for u in uncertainties:
        if u.linked_to is None:
            continue
        target = ids.get(u.linked_to)
        if target is None:
            problems.append(f"{u.where}: Linked_To names {u.linked_to!r}, which is not an X_ID")
        elif target.linked_to is not None:
            problems.append(f"{u.where}: Linked_To must name an X_ID that is not itself linked")
        elif target.x_id == u.x_id:
            problems.append(f"{u.where}: an uncertainty cannot be linked to itself")
    if uncertainties and samples == 0:
        problems.append("Setup: Uncertainty_Table has rows but Samples_Per_Future is 0")

    # Outputs
    input_variables, output_variables = [], []
    known = variable_catalogue()
    try:
        for number, row in _rows(wb["Outputs"], OUTPUT_COLUMNS):
            where = f"Outputs row {number}"
            if not _yes(row["Export"], where, "Export", problems, False):
                continue
            kind, variable = _text(row["Kind"]).capitalize(), _text(row["Variable"])
            if kind not in ("Input", "Output"):
                problems.append(f"{where}: Kind must be Input or Output")
                continue
            if variable not in known[kind]:
                other = "Output" if kind == "Input" else "Input"
                if variable in known[other]:
                    problems.append(f"{where}: {variable} is an {other.lower()} variable; set Kind to {other}")
                else:
                    hint = _closest(variable, known[kind])
                    problems.append(
                        f"{where}: {variable!r} is not an {kind.lower()} variable of the model"
                        + (f"; did you mean {hint}?" if hint else "")
                    )
                continue
            (input_variables if kind == "Input" else output_variables).append(variable)
    except InterfaceError as error:
        problems.extend(error.problems)
    wb.close()

    if problems:
        raise InterfaceError(problems)
    return Interface(
        path=path, run_name=run_name, base_scenario=base, samples_per_future=samples,
        random_seed=seed if seed is not None else 12345, run_central_future=central,
        apply_vre_ceilings=ceilings, model_first_year=first_year,
        model_last_year=last_year, futures=futures,
        changes=changes, uncertainties=uncertainties,
        input_variables=input_variables, output_variables=output_variables,
    )
