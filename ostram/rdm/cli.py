"""The ``python -m ostram rdm`` command.

Subcommands:

``template``
    Write a blank interface workbook, with dropdowns filled from the prepared
    root when one exists.
``generate``
    Validate an interface and write its run folder, without running anything.
``run``
    ``generate``, then build, compile and solve every future with the normal
    pipeline, then ``collect``.
``collect``
    Gather a run's solved futures into ``RDM_Inputs.csv`` and ``RDM_Outputs.csv``.

Every subcommand works on the active setup and workspace, so
``python -m ostram --setup unescap_training rdm run --interface my.xlsx``
runs the futures on the training model.
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path
import subprocess
import sys
from typing import Sequence

import yaml

from ostram.paths import resolve_paths
from ostram.pipeline.scenarios.registry import FUTURES_ENV, load_registry
from ostram.rdm.collect import clear_future_folders, collect_run
from ostram.rdm.design import write_run_folder
from ostram.rdm.interface import Interface, InterfaceError, read_interface, write_template

PROGRAM = "python -m ostram rdm"


def rdm_root() -> Path:
    """Return ``<active workspace>/rdm``, where every run folder lives."""

    return resolve_paths().workspace / "rdm"


def setup_has_vre_ceilings() -> bool:
    """Return whether the active setup's registry defines shared VRE ceilings."""

    raw = yaml.safe_load(resolve_paths().scenario_registry.read_text(encoding="utf-8")) or {}
    return bool(raw.get("vre_ceilings"))


def resolve_ceilings(choice: str) -> bool:
    """Turn the interface's Auto / Y / N into a yes or no for this setup."""

    available = setup_has_vre_ceilings()
    if choice == "Auto":
        return available
    if choice == "Y" and not available:
        raise ValueError(
            "Setup > Apply_VRE_Ceilings is Y, but this setup defines no shared VRE "
            "ceilings. Use Auto or N."
        )
    return choice == "Y"


def _root_lists(root: str) -> tuple[list[str], list[str], int | None, int | None]:
    """Return technologies, demand fuels and the year range of a prepared root, if any."""

    folder = resolve_paths().a1_outputs / f"A1_Outputs_{root}"
    technologies: set[str] = set()
    fuels: list[str] = []
    years: set[int] = set()
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl is a dependency
        return [], [], None, None
    params = folder / "A-O_Parametrization.xlsx"
    if params.is_file():
        wb = load_workbook(params, read_only=True)
        for sheet in ("Primary Techs", "Secondary Techs", "Demand Techs", "VariableCost"):
            if sheet not in wb.sheetnames:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
            for cell in header:
                try:
                    year = int(float(cell))
                except ValueError:
                    continue
                if 2000 <= year <= 2100:
                    years.add(year)
            if "Tech" in header:
                column = header.index("Tech")
                technologies.update(str(r[column]).strip() for r in rows if r[column])
        wb.close()
    demand = folder / "A-O_Demand.xlsx"
    if demand.is_file():
        wb = load_workbook(demand, read_only=True)
        if "Demand_Projection" in wb.sheetnames:
            rows = wb["Demand_Projection"].iter_rows(values_only=True)
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
            if "Fuel/Tech" in header:
                column = header.index("Fuel/Tech")
                fuels = [str(r[column]).strip() for r in rows if r[column]]
        wb.close()
    return (sorted(technologies), fuels,
            min(years) if years else None, max(years) if years else None)


def cmd_template(args: argparse.Namespace) -> int:
    registry = load_registry()
    roots = [r for r in registry.root_names if r not in registry.support_scenarios]
    roots += [r for r in registry.root_names if r in registry.support_scenarios]
    root = args.root or ("B_Optimised_VRE" if "B_Optimised_VRE" in roots else roots[0])
    technologies, fuels, first, last = _root_lists(root)
    path = write_template(
        args.out, roots=roots, base_scenario=root, technologies=technologies, fuels=fuels,
        model_first_year=first or 2023, model_last_year=last or 2050,
        examples=not args.blank,
    )
    print(f"Wrote the RDM interface template: {path}")
    if not technologies:
        print(
            f"The root {root} is not prepared in this workspace yet, so the selector "
            "dropdown is empty. Run the root once (for example with --compile-only) "
            "and write the template again to fill it."
        )
    return 0


def _generate(interface: Interface, *, overwrite: bool) -> tuple[Path, list[str]]:
    registry = load_registry()
    if interface.base_scenario not in registry.root_names:
        raise InterfaceError([
            f"Setup: Base_Scenario {interface.base_scenario!r} is not a root of this setup; "
            f"choose one of {', '.join(registry.root_names)}"
        ])
    run_dir = rdm_root() / interface.run_name
    specs = write_run_folder(
        interface, run_dir,
        apply_vre_ceilings=resolve_ceilings(interface.apply_vre_ceilings),
        reserved_names=[*registry.root_names, *registry.decision_scenarios],
        overwrite=overwrite,
    )
    names = [spec.name for spec in specs]
    print(f"Run folder: {run_dir}")
    print(f"{len(names)} future(s) from {len(interface.active_futures)} named future(s), "
          f"base scenario {interface.base_scenario}: {', '.join(names[:12])}"
          + (" ..." if len(names) > 12 else ""))
    return run_dir, names


def cmd_generate(args: argparse.Namespace) -> int:
    interface = read_interface(args.interface)
    _generate(interface, overwrite=args.overwrite)
    return 0


def _run_pipeline(futures_file: Path, names: Sequence[str], args: argparse.Namespace) -> int:
    """Run ``python -m ostram run`` on the futures in-process, with the futures registered."""

    from ostram.pipeline import orchestration

    argv = ["python -m ostram run", "--skip-pull", "--scenarios", ",".join(names)]
    if args.compile_only:
        argv.append("--compile-only")
    if args.env_name:
        argv += ["--env-name", args.env_name]
    if args.a_result_seed:
        argv += ["--a-result-seed", args.a_result_seed]
    if args.verbose:
        argv.append("--verbose")
    previous_argv, previous_env = sys.argv, os.environ.get(FUTURES_ENV)
    sys.argv = argv
    os.environ[FUTURES_ENV] = str(futures_file)
    try:
        orchestration.main()
    except subprocess.CalledProcessError as error:
        print(f"\nThe pipeline stopped (exit {error.returncode}).", file=sys.stderr)
        return error.returncode or 1
    except Exception as error:  # the run reporter has already logged the detail
        print(f"\nThe pipeline stopped: {error}", file=sys.stderr)
        return 1
    finally:
        sys.argv = previous_argv
        if previous_env is None:
            os.environ.pop(FUTURES_ENV, None)
        else:
            os.environ[FUTURES_ENV] = previous_env
    return 0


def _report(summary: dict) -> None:
    print(f"Solved: {len(summary['solved'])}, compiled only: {len(summary['compiled'])}, "
          f"missing: {len(summary['missing'])}")
    if summary["missing"]:
        print(f"Missing futures: {', '.join(summary['missing'])}")
    if summary.get("not_found"):
        print("Ticked on the Outputs sheet but empty in every future (the model gives them no "
              f"values), so not in the CSVs: {', '.join(summary['not_found'])}")
    for label, key in (("RDM inputs", "inputs"), ("RDM outputs", "outputs"), ("Run status", "status")):
        if summary.get(key):
            print(f"{label}: {summary[key]}")


def cmd_run(args: argparse.Namespace) -> int:
    interface = read_interface(args.interface)
    run_dir, names = _generate(interface, overwrite=args.overwrite)
    cleared = clear_future_folders(names)
    if cleared:
        print(f"Cleared {len(cleared)} result folder(s) left by an earlier run of the same Future_IDs.")
    code = _run_pipeline(run_dir / "futures.yaml", names, args)
    summary = collect_run(run_dir)
    _report(summary)
    return code


def cmd_collect(args: argparse.Namespace) -> int:
    run = Path(args.run)
    run_dir = run if run.is_dir() else rdm_root() / args.run
    if not (run_dir / "rdm_manifest.json").is_file():
        raise FileNotFoundError(f"no RDM run folder at {run_dir}")
    _report(collect_run(run_dir))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog=PROGRAM, description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="action", required=True)

    template = sub.add_parser("template", help="Write a blank interface workbook.")
    template.add_argument("--out", default="Interface_OSTRAM_RDM.xlsx",
                          help="Workbook to write (default: Interface_OSTRAM_RDM.xlsx).")
    template.add_argument("--root", help="Root whose technologies fill the dropdowns.")
    template.add_argument("--blank", action="store_true", help="Leave out the worked example.")
    template.set_defaults(func=cmd_template)

    generate = sub.add_parser("generate", help="Validate an interface and write its run folder only.")
    generate.add_argument("--interface", required=True)
    generate.add_argument("--overwrite", action="store_true", help="Replace an existing run folder.")
    generate.set_defaults(func=cmd_generate)

    run = sub.add_parser("run", help="Generate, run every future, and collect the RDM CSVs.")
    run.add_argument("--interface", required=True)
    run.add_argument("--overwrite", action="store_true", help="Replace an existing run folder.")
    run.add_argument("--compile-only", action="store_true",
                     help="Build and compile the futures but run no solver.")
    run.add_argument("--env-name", help="Conda environment, as for `ostram run`.")
    run.add_argument("--a-result-seed", help="Solved A_Calibrated_BAU result, for C_Target_VRE futures.")
    run.add_argument("--verbose", action="store_true")
    run.set_defaults(func=cmd_run)

    collect = sub.add_parser("collect", help="Re-collect a run's CSVs from the solved futures.")
    collect.add_argument("--run", required=True, help="Run_Name, or the path of a run folder.")
    collect.set_defaults(func=cmd_collect)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(sys.argv[1:] if argv is None else argv)
    try:
        return args.func(args)
    except (InterfaceError, FileNotFoundError, FileExistsError, ValueError) as error:
        print(f"\n{error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
