"""RDM futures: the Excel interface, the future expansion, the registry overlay,
the demand edits and the collection of results.  No solver runs."""

from __future__ import annotations

import csv
import json
import os
import sys
from pathlib import Path
import tempfile
import unittest
from unittest import mock

import numpy as np
from openpyxl import Workbook, load_workbook
import pandas as pd
import yaml

from ostram.pipeline.scenarios import registry as registry_module
from ostram.pipeline.scenarios import scenario_edits as edits
from ostram.rdm import collect, design, interface
from ostram.rdm.interface import Change, InterfaceError

ROOT = Path(__file__).resolve().parents[2]
TRAINING_REGISTRY = ROOT / "examples" / "unescap_training" / "config" / "scenarios" / "registry.json"
YEARS = list(range(2023, 2051))


def _template(folder: Path, cells: dict | None = None) -> Path:
    """Write the example template, change ``{(sheet, cell): value}`` and return its path."""

    path = folder / "interface.xlsx"
    interface.write_template(path, technologies=["PWRSPVBGDXX", "MINCOABGD"], fuels=["ELCBGDXX03"])
    if cells:
        wb = load_workbook(path)
        for (sheet, cell), value in cells.items():
            wb[sheet][cell] = value
        wb.save(path)
    return path


class InterfaceTemplateTests(unittest.TestCase):
    def test_template_round_trips_with_its_example(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec = interface.read_interface(_template(Path(tmp), {("Setup", "B4"): 3}))
        self.assertEqual(spec.base_scenario, "B_Optimised_VRE")
        self.assertEqual([f.future_id for f in spec.futures], ["F00", "F01", "F02"])
        self.assertEqual(spec.changes["F00"], [])
        self.assertEqual(len(spec.changes["F02"]), 2)
        self.assertEqual([u.x_id for u in spec.uncertainties], ["X1", "X2"])
        self.assertIn("TotalCapacityAnnual", spec.output_variables)
        self.assertNotIn("RateOfActivity", spec.output_variables)
        self.assertIn("CapitalCost", spec.input_variables)

    def test_uncertainties_without_samples_are_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(InterfaceError) as caught:
                interface.read_interface(_template(Path(tmp)))
        self.assertIn("Samples_Per_Future is 0", str(caught.exception))

    def test_every_problem_is_reported_at_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp), {
                ("Setup", "B4"): 2,
                ("Futures", "A3"): "1bad",                    # bad Future_ID
                ("Future_Changes", "B2"): "Nowhere",          # unknown sheet
                ("Future_Changes", "D4"): "Fuel",             # fuel selector on a tech sheet
                ("Future_Changes", "I3"): None,               # fine: open-ended Multiply
                ("Uncertainty_Table", "C3"): "F09",           # unknown future
                ("Uncertainty_Table", "M2"): "X9",            # unknown link
                ("Uncertainty_Table", "J3"): 0.1,             # Max below Min
            })
            with self.assertRaises(InterfaceError) as caught:
                interface.read_interface(path)
        message = str(caught.exception)
        for fragment in ("Futures row 3", "Sheet 'Nowhere'", "Selector_Type is Tech",
                         "'F09'", "'X9'", "Max_Value must not be below Min_Value"):
            self.assertIn(fragment, message)

    def test_outputs_sheet_lists_every_model_variable_with_its_sets(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp))
            rows = list(load_workbook(path)["Outputs"].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Kind", "Variable", "Export", "Indexed_By"))
        listed = {(r[0], r[1]): (r[2], r[3]) for r in rows[1:]}
        catalogue = interface.variable_catalogue()
        expected = {(k, v) for k in ("Input", "Output") for v in catalogue[k]}
        self.assertEqual(set(listed), expected)
        self.assertGreaterEqual(len(catalogue["Input"]), 50)
        self.assertEqual(listed[("Input", "AccumulatedAnnualDemand")], ("N", "REGION, FUEL, YEAR"))
        self.assertEqual(listed[("Input", "CapitalCost")][0], "Y")
        self.assertIn("EMISSION", listed[("Input", "AnnualEmissionLimit")][1])

    def test_ticking_a_variable_exports_it_and_typos_are_caught(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp), {("Setup", "B4"): 1})
            wb = load_workbook(path)
            ws = wb["Outputs"]
            row_of = {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1)}
            ws.cell(row_of["AccumulatedAnnualDemand"], 3, "Y")
            wb.save(path)
            spec = interface.read_interface(path)
            self.assertIn("AccumulatedAnnualDemand", spec.input_variables)

            ws.cell(row_of["InputActivityRatio"], 2, "inputactivityratio")
            ws.cell(row_of["InputActivityRatio"], 3, "Y")
            ws.cell(row_of["TotalCapacityAnnual"], 1, "Input")
            wb.save(path)
            with self.assertRaises(InterfaceError) as caught:
                interface.read_interface(path)
        message = str(caught.exception)
        self.assertIn("did you mean InputActivityRatio?", message)
        self.assertIn("TotalCapacityAnnual is an output variable; set Kind to Output", message)

    def test_uncertainty_table_uses_the_osemosys_rdm_column_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp), {("Setup", "B4"): 2})
            header = [c.value for c in load_workbook(path)["Uncertainty_Table"][1]]
            self.assertEqual(header[:3], ["X_Num", "X_Category", "Involved_Scenarios"])
            for name in ("Exact_Parameters_Involved_in_Osemosys", "Involved_First_Sets_in_Osemosys",
                         "X_Mathematical_Type", "Min_Value", "Max_Value", "Initial_Year_of_Uncertainty",
                         "X_Plain_English_Description"):
                self.assertIn(name, header)
            new = interface.read_interface(path)
            # a workbook with the earlier short headers reads the same
            wb = load_workbook(path)
            for cell, short in zip(wb["Uncertainty_Table"][1], interface.UNCERTAINTY_COLUMNS):
                cell.value = short
            wb.save(path)
            old = interface.read_interface(path)
        self.assertEqual(new.uncertainties, [u for u in old.uncertainties])

    def test_x_num_may_be_a_number_and_lists_may_use_semicolons(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp), {
                ("Setup", "B4"): 2,
                ("Uncertainty_Table", "A2"): 1, ("Uncertainty_Table", "A3"): 2,
                ("Uncertainty_Table", "C3"): "F01 ; F02",
                ("Uncertainty_Table", "F3"): "Techs", ("Uncertainty_Table", "G3"): "MINGASBGD ; MINGASIND",
                ("Uncertainty_Table", "M3"): 1,
            })
            spec = interface.read_interface(path)
        x1, x2 = spec.uncertainties
        self.assertEqual((x1.x_id, x2.x_id), ("X1", "X2"))
        self.assertEqual(x2.applies_to, ("F01", "F02"))
        self.assertEqual(x2.selector, "MINGASBGD, MINGASIND")
        self.assertEqual(x2.linked_to, "X1")

    def test_ramp_needs_both_years(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = _template(Path(tmp), {("Setup", "B4"): 1, ("Future_Changes", "H2"): None})
            with self.assertRaises(InterfaceError) as caught:
                interface.read_interface(path)
        self.assertIn("Multiply_Ramp needs both Year_From and Year_To", str(caught.exception))


class ChangeTranslationTests(unittest.TestCase):
    def _change(self, operation, value, year_from=None, year_to=None, selector_type="Tech_Prefix"):
        return Change("Secondary Techs", "CapitalCost", selector_type, "PWRSPV", operation,
                      value, year_from, year_to, "basis", "row 2")

    def test_multiply_with_and_without_years(self):
        whole = design.change_to_edits(self._change("Multiply", 1.1), 2023, 2050)
        self.assertEqual(whole, [{"sheet": "Secondary Techs", "param": "CapitalCost",
                                  "tech_prefix": "PWRSPV", "note": "basis [row 2]",
                                  "op": "multiply", "factor": 1.1}])
        part = design.change_to_edits(self._change("Multiply", 1.1, 2040), 2023, 2050)
        self.assertEqual(part[0]["years"], list(range(2040, 2051)))

    def test_set_flat_and_set_values(self):
        flat = design.change_to_edits(self._change("Set", 5), 2023, 2050)[0]
        self.assertEqual((flat["op"], flat["value"]), ("set_flat", 5.0))
        values = design.change_to_edits(self._change("Set", 5, 2030, 2031), 2023, 2050)[0]
        self.assertEqual(values["values"], {"2030": 5.0, "2031": 5.0})

    def test_ramp_rises_linearly_then_holds(self):
        result = design.change_to_edits(self._change("Multiply_Ramp", 0.5, 2030, 2035), 2023, 2050)
        factors = [(e["years"][0], e["factor"]) for e in result[:-1]]
        self.assertEqual(factors, [(2031, 0.9), (2032, 0.8), (2033, 0.7), (2034, 0.6), (2035, 0.5)])
        self.assertEqual(result[-1]["years"], list(range(2036, 2051)))
        self.assertEqual(result[-1]["factor"], 0.5)

    def test_selectors(self):
        techs = design.change_to_edits(self._change("Multiply", 2, selector_type="Techs"), 2023, 2050)[0]
        self.assertEqual(techs["techs"], ["PWRSPV"])
        copy = design.change_to_edits(
            Change("VariableCost", "VariableCost", "Tech", "MINCOABGD", "Copy_From_Tech", "MINCOAIND"),
            2023, 2050,
        )[0]
        self.assertEqual(copy, {"sheet": "VariableCost", "param": "VariableCost", "tech": "MINCOABGD",
                                "op": "copy_from_tech", "source_tech": "MINCOAIND"})


class ExpansionTests(unittest.TestCase):
    def test_latin_hypercube_fills_every_stratum_once(self):
        draws = design.latin_hypercube(20, 3, np.random.default_rng(1))
        for column in draws.T:
            self.assertEqual(sorted((column * 20).astype(int)), list(range(20)))

    def _spec(self, tmp, cells=None):
        return interface.read_interface(_template(Path(tmp), cells))

    def test_layered_samples_share_draws_across_futures(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec = self._spec(tmp, {("Setup", "B4"): 4})
        futures = design.expand_futures(spec)
        names = [f.name for f in futures]
        self.assertEqual(names[:5], ["F00", "F00_s001", "F00_s002", "F00_s003", "F00_s004"])
        self.assertEqual(len(futures), 15)
        by_name = {f.name: f for f in futures}
        self.assertEqual(by_name["F01_s003"].sampled, by_name["F02_s003"].sampled)
        self.assertNotEqual(by_name["F01_s003"].sampled, by_name["F01_s004"].sampled)
        # The future's own changes come first, then the sampled uncertainties.
        own = by_name["F02"].edits
        self.assertEqual(by_name["F02_s001"].edits[: len(own)], own)
        self.assertGreater(len(by_name["F02_s001"].edits), len(own))
        for value in (by_name[f"F00_s00{i}"].sampled["X2"] for i in range(1, 5)):
            self.assertTrue(0.8 <= value <= 1.5)

    def test_central_future_can_be_left_out_and_seed_is_reproducible(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec = self._spec(tmp, {("Setup", "B4"): 2, ("Setup", "B6"): "N"})
        first = design.expand_futures(spec)
        self.assertNotIn("F01", [f.name for f in first])
        self.assertEqual([f.sampled for f in first], [f.sampled for f in design.expand_futures(spec)])

    def test_applies_to_and_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec = self._spec(tmp, {
                ("Setup", "B4"): 5,
                ("Uncertainty_Table", "C2"): "F01",
                ("Uncertainty_Table", "C3"): "F01",
                ("Uncertainty_Table", "M3"): "X1",
            })
        futures = {f.name: f for f in design.expand_futures(spec)}
        self.assertIn("F00", futures)
        self.assertNotIn("F00_s001", futures)  # no uncertainty applies to F00
        for s in range(1, 6):
            draw = futures[f"F01_s00{s}"].sampled
            # X1 spans 0.9-1.1 and X2 0.8-1.5: linked, so the same position in each range.
            self.assertAlmostEqual((draw["X1"] - 0.9) / 0.2, (draw["X2"] - 0.8) / 0.7, places=4)


class RunFolderAndRegistryTests(unittest.TestCase):
    def test_run_folder_registers_its_futures(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            spec = interface.read_interface(_template(tmp, {("Setup", "B4"): 1}))
            run_dir = tmp / "rdm" / spec.run_name
            futures = design.write_run_folder(spec, run_dir, apply_vre_ceilings=False)
            self.assertEqual([f.name for f in futures],
                             ["F00", "F00_s001", "F01", "F01_s001", "F02", "F02_s001"])
            document = yaml.safe_load((run_dir / "scenarios" / "F01" / "scenario.yaml").read_text())
            self.assertEqual(document["base_scenario"], "B_Optimised_VRE")
            self.assertFalse(document["apply_vre_ceiling_layer"])
            with (run_dir / "RDM_Experiment_Design.csv").open() as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(rows[1]["Future_ID"], "F00_s001")
            self.assertNotEqual(rows[1]["X1"], "")

            with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: str(run_dir / "futures.yaml")}):
                loaded = registry_module.load_registry(TRAINING_REGISTRY)
            self.assertIn("F01_s001", loaded.decision_scenarios)
            self.assertEqual(loaded.derived_by_name["F01"].config_root, run_dir / "scenarios")
            self.assertEqual(loaded.select("F02,F00"), ("F00", "F02"))
            self.assertEqual(loaded.required_roots(["F01"]), ("B_Optimised_VRE",))

            with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: ""}):
                self.assertNotIn("F01", registry_module.load_registry(TRAINING_REGISTRY).decision_scenarios)

            with self.assertRaises(FileExistsError):
                design.write_run_folder(spec, run_dir, apply_vre_ceilings=False)
            design.write_run_folder(spec, run_dir, apply_vre_ceilings=False, overwrite=True)

    def test_reserved_names_and_non_root_bases_are_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            spec = interface.read_interface(_template(tmp, {("Setup", "B4"): 1}))
            with self.assertRaisesRegex(ValueError, "already registered"):
                design.write_run_folder(spec, tmp / "a", apply_vre_ceilings=False, reserved_names=["F01"])
            spec.base_scenario = "Not_A_Root"
            design.write_run_folder(spec, tmp / "b", apply_vre_ceilings=False)
            with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: str(tmp / "b" / "futures.yaml")}):
                with self.assertRaisesRegex(ValueError, "non-root base"):
                    registry_module.load_registry(TRAINING_REGISTRY)


def _demand_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Demand_Projection"
    ws.append(["Demand/Share", "Fuel/Tech", "Name", "Projection.Mode", "Projection.Parameter",
               *[str(y) for y in YEARS]])
    ws.append(["Demand", "ELCBGDXX03", "BGD", "User defined", 0, *[100.0] * len(YEARS)])
    ws.append(["Demand", "ELCINDEA03", "IND", "", 0, *[200.0] * len(YEARS)])
    return wb


def _log():
    return {"scenario": "F01", "cells": [], "rows_created": [], "skipped": []}


class DemandEditTests(unittest.TestCase):
    def _values(self, ws, row):
        return [ws.cell(row=row, column=c).value for c in range(6, 6 + len(YEARS))]

    def test_multiply_values_and_set_flat(self):
        wb = _demand_workbook()
        ws = wb["Demand_Projection"]
        log = _log()
        edits.apply_demand_edit(ws, {"sheet": "Demand_Projection", "fuel_prefix": "ELC",
                                     "op": "multiply", "factor": 1.5, "years": [2050]}, log)
        self.assertEqual(self._values(ws, 2)[-1], 150.0)
        self.assertEqual(self._values(ws, 2)[0], 100.0)
        self.assertEqual(self._values(ws, 3)[-1], 300.0)
        self.assertEqual(ws.cell(row=3, column=4).value, "User defined")
        edits.apply_demand_edit(ws, {"sheet": "Demand_Projection", "fuel": "ELCBGDXX03",
                                     "values": {"2023": 7}}, log)
        self.assertEqual(self._values(ws, 2)[0], 7.0)
        edits.apply_demand_edit(ws, {"sheet": "Demand_Projection", "fuels": ["ELCINDEA03"],
                                     "op": "set_flat", "value": 1}, log)
        self.assertEqual(set(self._values(ws, 3)), {1.0})
        self.assertTrue(all(cell["workbook"] == "A-O_Demand.xlsx" for cell in log["cells"]))

    def test_edits_that_match_nothing_stop(self):
        ws = _demand_workbook()["Demand_Projection"]
        with self.assertRaisesRegex(edits.ScenarioEditError, "Did you mean ELCBGDXX03"):
            edits.apply_demand_edit(ws, {"fuel": "ELCBGDXX3", "op": "set_flat", "value": 1}, _log())
        with self.assertRaisesRegex(edits.ScenarioEditError, "No demand fuel starts"):
            edits.apply_demand_edit(ws, {"fuel_prefix": "GAS", "op": "set_flat", "value": 1}, _log())
        with self.assertRaisesRegex(edits.ScenarioEditError, "2070"):
            edits.apply_demand_edit(ws, {"fuel": "ELCBGDXX03", "op": "multiply",
                                         "factor": 2, "years": [2070]}, _log())
        with self.assertRaisesRegex(ValueError, "holds only"):
            edits.apply_demand_edit(ws, {"fuel": "ELCBGDXX03", "param": "CapitalCost",
                                         "op": "set_flat", "value": 1}, _log())

    def test_build_scenario_edits_demand_only_when_asked(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            root = tmp / "A1" / "A1_Outputs_B_Optimised_VRE"
            root.mkdir(parents=True)
            params = Workbook()
            sheet = params.active
            sheet.title = "Secondary Techs"
            sheet.append(["Tech", "Parameter", "Projection.Mode", *YEARS])
            sheet.append(["PWRSPVBGDXX", "CapitalCost", "User defined", *[10.0] * len(YEARS)])
            params.save(root / edits.PARAM_FILE)
            _demand_workbook().save(root / edits.DEMAND_FILE)
            configs = tmp / "scenarios"
            for name, scenario_edits in (
                ("F01", [{"sheet": "Demand_Projection", "fuel": "ELCBGDXX03", "op": "multiply", "factor": 2}]),
                ("F02", [{"sheet": "Secondary Techs", "tech": "PWRSPVBGDXX", "param": "CapitalCost",
                          "op": "multiply", "factor": 0.5}]),
            ):
                (configs / name).mkdir(parents=True)
                (configs / name / "scenario.yaml").write_text(yaml.safe_dump({
                    "scenario": name, "base_scenario": "B_Optimised_VRE",
                    "apply_vre_ceiling_layer": False, "edits": scenario_edits,
                }))
            with mock.patch("ostram.pipeline.scenarios.rules.apply_base_year_pin.calibration_root",
                            return_value=None):
                edits.build_scenario("F01", a1_outputs=tmp / "A1", configs=configs)
                edits.build_scenario("F02", a1_outputs=tmp / "A1", configs=configs)
            demand = load_workbook(tmp / "A1" / "A1_Outputs_F01" / edits.DEMAND_FILE)["Demand_Projection"]
            self.assertEqual(demand.cell(row=2, column=6).value, 200.0)
            source = (root / edits.DEMAND_FILE).read_bytes()
            self.assertEqual((tmp / "A1" / "A1_Outputs_F02" / edits.DEMAND_FILE).read_bytes(), source)


class CollectTests(unittest.TestCase):
    def test_collect_keeps_selected_variables_and_reports_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            spec = interface.read_interface(_template(tmp, {("Setup", "B4"): 1}))
            run_dir = tmp / "rdm" / spec.run_name
            design.write_run_folder(spec, run_dir, apply_vre_ceilings=False)
            execution = tmp / "execution"
            execution.mkdir()
            (execution / "Config_MOMF_T1_AB.yaml").write_text(yaml.safe_dump({
                "executables": "Executables", "preprocess_data_name": "Pre_processed_",
                "output_files": "_output",
            }))
            for name, solved in (("F00", True), ("F01", False)):
                folder = execution / "Executables" / f"{name}_0"
                folder.mkdir(parents=True)
                pd.DataFrame({"YEAR": [2030, 2031], "TECHNOLOGY": ["PWRSPVBGDXX"] * 2,
                              "CapitalCost": [1.0, 2.0], "CapacityFactor": [None, None],
                              "Junk": [9, 9]}).to_csv(folder / f"{name}_0_Input.csv", index=False)
                if solved:
                    pd.DataFrame({"YEAR": [2030], "TECHNOLOGY": ["PWRSPVBGDXX"],
                                  "TotalCapacityAnnual": [3.0], "RateOfActivity": [1.0]}).to_csv(
                        folder / f"Pre_processed_{name}_0_output.csv", index=False)
            summary = collect.collect_run(run_dir, execution)
            self.assertEqual(summary["solved"], ["F00"])
            self.assertEqual(summary["compiled"], ["F01"])
            self.assertIn("F02_s001", summary["missing"])
            inputs = pd.read_csv(summary["inputs"])
            self.assertEqual(list(inputs.columns[:3]), ["Future_ID", "Named_Future", "Sample"])
            self.assertIn("CapitalCost", inputs.columns)
            self.assertNotIn("Junk", inputs.columns)
            self.assertEqual(sorted(inputs["Future_ID"].unique()), ["F00", "F01"])
            outputs = pd.read_csv(summary["outputs"])
            self.assertEqual(list(outputs["TotalCapacityAnnual"]), [3.0])
            self.assertNotIn("RateOfActivity", outputs.columns)
            # ticked by default but absent from every future's CSV
            self.assertIn("SpecifiedAnnualDemand", summary["not_found"])
            self.assertNotIn("CapitalCost", summary["not_found"])
            self.assertNotIn("TotalCapacityAnnual", summary["not_found"])

    def test_clear_removes_only_the_named_future_folders(self):
        with tempfile.TemporaryDirectory() as tmp:
            execution = Path(tmp)
            (execution / "Config_MOMF_T1_AB.yaml").write_text(yaml.safe_dump({
                "executables": "Executables", "preprocess_data_name": "Pre_processed_",
                "output_files": "_output",
            }))
            for name in ("F01_0", "F02_0", "B_Optimised_VRE_0"):
                (execution / "Executables" / name).mkdir(parents=True)
                (execution / "Executables" / name / "old_output.csv").write_text("x")
            removed = collect.clear_future_folders(["F01", "F09"], execution)
            self.assertEqual([p.name for p in removed], ["F01_0"])
            self.assertFalse((execution / "Executables" / "F01_0").exists())
            self.assertTrue((execution / "Executables" / "F02_0").exists())
            self.assertTrue((execution / "Executables" / "B_Optimised_VRE_0").exists())


class SolverBehaviourTests(unittest.TestCase):
    """B2 treats RDM futures like their root, and keeps going past a failed one."""

    def _futures_env(self, tmp: Path) -> dict:
        spec = interface.read_interface(_template(tmp, {("Setup", "B4"): 1}))
        run_dir = tmp / "rdm" / spec.run_name
        design.write_run_folder(spec, run_dir, apply_vre_ceilings=False)
        return {registry_module.FUTURES_ENV: str(run_dir / "futures.yaml")}

    def test_futures_use_their_roots_solver_settings(self):
        from ostram.pipeline.execution import orchestrator

        with tempfile.TemporaryDirectory() as tmp:
            env = self._futures_env(Path(tmp))
            with mock.patch.dict(os.environ, env), \
                    mock.patch("ostram.pipeline.scenarios.registry.resolve_paths") as paths:
                paths.return_value.scenario_registry = TRAINING_REGISTRY
                self.assertEqual(orchestrator.solver_settings_scenario("F01"), "B_Optimised_VRE")
                self.assertEqual(orchestrator.solver_settings_scenario("F01_s001"), "B_Optimised_VRE")
                self.assertEqual(orchestrator.solver_settings_scenario("B_Optimised_VRE"), "B_Optimised_VRE")
        with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: ""}):
            self.assertEqual(orchestrator.solver_settings_scenario("F01"), "F01")

    def _cbc_command(self, scenario, future_root=None):
        """Return the CBC command B2 prepares for ``scenario``."""

        from ostram.pipeline.execution import orchestrator
        from tests.regression.test_b2_orchestration import _base_params

        class Stop(Exception):
            pass

        dependencies = orchestrator.ScenarioExecutionDependencies(
            run_process=lambda *a, **k: None, check_environment=lambda solver: None,
            get_executable=lambda name: name, path_exists=lambda path: False,
            remove_file=lambda path: None, python_executable=sys.executable,
        )
        real = orchestrator.SolverAdapter(dependencies)
        captured = {}

        class Recorder:
            def prepare_command(self, solver, params, paths):
                captured["command"] = real.prepare_command(solver, params, paths)
                raise Stop

        params = _base_params(solver="cbc", execute_model=True, create_matrix=True,
                              cbc_primal_tolerance="1e-10",
                              cbc_primal_tolerance_by_scenario={"B_Optimised_VRE": "1e-9"})
        with mock.patch.object(orchestrator, "solver_settings_scenario",
                               return_value=future_root or scenario):
            with self.assertRaises(Stop):
                orchestrator.execute_scenario(params, scenario, Path("."), dependencies,
                                              solver_adapter=Recorder(),
                                              matrix_runner=lambda *a: None)
        return captured["command"]

    def test_futures_get_their_roots_tolerance_and_primal_simplex(self):
        future = self._cbc_command("F01", future_root="B_Optimised_VRE")
        self.assertEqual(future[future.index("primalTolerance") + 1], "1e-9")
        self.assertIn("-primalS", future)
        self.assertLess(future.index("-primalS"), future.index("solve"))
        root = self._cbc_command("B_Optimised_VRE")
        self.assertEqual(root[root.index("primalTolerance") + 1], "1e-9")
        self.assertNotIn("-primalS", root)
        other = self._cbc_command("A_Calibrated_BAU")
        self.assertEqual(other[other.index("primalTolerance") + 1], "1e-10")
        self.assertNotIn("-primalS", other)

    def _plan(self, scenarios):
        from ostram.pipeline.execution import orchestrator

        return orchestrator.B2RunPlan(
            here=Path("."), params={"execute_model": True, "create_matrix": True, "parallel": False},
            params_a2={}, scenarios=scenarios, main_scenario_name="", base_input_path="",
            template_path="", base_output_path="",
        )

    def _dependencies(self, solved):
        def main_executer(params, scenario, here):
            if scenario == "F01":
                raise RuntimeError("CBC did not produce an optimal solution")
            solved.append(scenario)
        return mock.Mock(main_executer=main_executer)

    def test_a_failed_future_does_not_stop_the_others(self):
        from ostram.pipeline.execution import orchestrator

        solved: list[str] = []
        with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: "futures.yaml"}):
            with self.assertRaisesRegex(RuntimeError, "1 future\\(s\\) failed in B2: F01"):
                orchestrator.run_execution_stage(self._plan(["F00", "F01", "F02"]), self._dependencies(solved))
        self.assertEqual(solved, ["F00", "F02"])

    def test_other_runs_still_stop_at_the_first_failure(self):
        from ostram.pipeline.execution import orchestrator

        solved: list[str] = []
        with mock.patch.dict(os.environ, {registry_module.FUTURES_ENV: ""}):
            with self.assertRaisesRegex(RuntimeError, "CBC did not produce"):
                orchestrator.run_execution_stage(self._plan(["F00", "F01", "F02"]), self._dependencies(solved))
        self.assertEqual(solved, ["F00"])


if __name__ == "__main__":
    unittest.main()
