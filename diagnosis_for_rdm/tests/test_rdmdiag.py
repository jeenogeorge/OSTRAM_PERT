"""Build a small RDM run folder and check every diagnosis step on it.

Run from the clone root:  python -m unittest discover -s diagnosis_for_rdm/tests
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))

from rdmdiag import allsets, infeasibility, perturbed  # noqa: E402
from rdmdiag.run import load_run, resolve_run  # noqa: E402
import run_diagnosis  # noqa: E402

YEARS = list(range(2023, 2031))
FUTURES = ["F00", "F00_s001", "F00_s002", "F00_s003", "F00_s004"]
DRAWS = {"F00_s001": (1.05, 1.2), "F00_s002": (0.95, 1.5), "F00_s003": (1.02, 1.1), "F00_s004": (0.98, 2.9)}
UNSOLVED = {"F00_s003": None, "F00_s004": "Infeasible - objective value 0"}   # None: no .sol at all


def _interface(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Setup"
    for row in [("Setting", "Value"), ("Run_Name", "test01"), ("Base_Scenario", "B_Optimised_VRE")]:
        ws.append(row)
    wb.create_sheet("Futures").append(("Future_ID", "Name", "Description", "Run"))
    wb["Futures"].append(("F00", "Baseline", "", "Y"))
    wb.create_sheet("Future_Changes").append(
        ("Future_ID", "Sheet", "Parameter", "Selector_Type", "Selector", "Operation", "Value",
         "Year_From", "Year_To", "Note"))
    ut = wb.create_sheet("Uncertainty_Table")
    ut.append(("X_ID", "Name", "Applies_To", "Sheet", "Parameter", "Selector_Type", "Selector",
               "Method", "Min", "Max", "Year_From", "Year_To", "Linked_To", "Note"))
    ut.append(("X1", "Demand", "ALL", "Demand_Projection", "SpecifiedAnnualDemand", "Fuel_Prefix", "ELC",
               "Multiply", 0.9, 1.1, None, None, None, ""))
    ut.append(("X2", "Coal min build", "ALL", "Secondary Techs", "TotalAnnualMinCapacityInvestment", "Tech",
               "PWRCOA", "Multiply", 1.0, 3.0, None, None, None, ""))
    ut.append(("X3", "Gas price", "ALL", "VariableCost", "VariableCost", "Tech", "MINGAS",
               "Multiply", 0.8, 1.5, None, None, None, ""))
    wb.save(path)


def _future_inputs(future: str) -> pd.DataFrame:
    x1, x2 = DRAWS.get(future, (1.0, 1.0))
    x3 = {"F00_s001": 0.9, "F00_s002": 1.3, "F00_s003": 1.1, "F00_s004": 1.4}.get(future, 1.0)
    rows = []
    for y in YEARS:
        rows += [
            {"REGION": "R1", "YEAR": y, "FUEL": "ELC01", "SpecifiedAnnualDemand": 100 * x1},
            {"REGION": "R1", "YEAR": y, "TECHNOLOGY": "PWRCOA", "TotalAnnualMinCapacityInvestment": 0.2 * x2},
            {"REGION": "R1", "YEAR": y, "TECHNOLOGY": "PWRCOA", "TotalAnnualMaxCapacityInvestment": 0.5},
            {"REGION": "R1", "YEAR": y, "TECHNOLOGY": "MINGAS", "MODE_OF_OPERATION": 1, "VariableCost": 3 * x3},
            {"REGION": "R1", "YEAR": y, "TECHNOLOGY": "PWRCOA", "FUEL": "COA", "MODE_OF_OPERATION": 1,
             "InputActivityRatio": 2.5},
            {"REGION": "R1", "YEAR": y, "EMISSION": "CO2", "AnnualEmissionLimit": 9999},
        ]
    return pd.DataFrame(rows)


def build_run(workspace: Path) -> Path:
    run_dir = workspace / "setups" / "unescap_training" / "rdm" / "test01"
    run_dir.mkdir(parents=True)
    _interface(run_dir / "interface.xlsx")
    (run_dir / "rdm_manifest.json").write_text(json.dumps({"futures": FUTURES}))
    pd.DataFrame([{"Future_ID": f, "Named_Future": "F00", "Sample": i,
                   "X1": DRAWS.get(f, ("", ""))[0], "X2": DRAWS.get(f, ("", ""))[1],
                   "X3": {"F00_s001": 0.9, "F00_s002": 1.3, "F00_s003": 1.1, "F00_s004": 1.4}.get(f, "")}
                  for i, f in enumerate(FUTURES)]).to_csv(run_dir / "RDM_Experiment_Design.csv", index=False)
    exe = workspace / "setups" / "unescap_training" / "execution" / "Executables"
    status, exported = [], []
    for i, f in enumerate(FUTURES):
        folder = exe / f"{f}_0"
        folder.mkdir(parents=True)
        inp = _future_inputs(f)
        inp.to_csv(folder / f"{f}_0_Input.csv", index=False)
        exported.append(inp.drop(columns=["InputActivityRatio", "AnnualEmissionLimit"])
                        .dropna(how="all", subset=["SpecifiedAnnualDemand", "TotalAnnualMinCapacityInvestment",
                                                   "TotalAnnualMaxCapacityInvestment", "VariableCost"])
                        .assign(Future_ID=f, Named_Future="F00", Sample=i))
        solved = f not in UNSOLVED
        if solved:
            (folder / f"{f}_0_Output.sol").write_text("Optimal - objective value 1\n")
            pd.DataFrame({"REGION": "R1", "YEAR": YEARS, "TotalDiscountedCost": [10.0 + i] * len(YEARS)}) \
                .to_csv(folder / f"{f}_0_Output.csv", index=False)
        elif UNSOLVED[f]:
            (folder / f"{f}_0_Output.sol").write_text(UNSOLVED[f] + "\n")
        status.append({"Future_ID": f, "Named_Future": "F00", "Sample": i,
                       "Status": "solved" if solved else "compiled",
                       "Input_CSV": str(folder / f"{f}_0_Input.csv"),
                       "Output_CSV": str(folder / f"{f}_0_Output.csv") if solved else ""})
    pd.DataFrame(status).to_csv(run_dir / "RDM_Run_Status.csv", index=False)
    pd.concat(exported).to_csv(run_dir / "RDM_Inputs.csv", index=False)
    return run_dir


class DiagnosisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ws = Path(cls.tmp.name)
        cls.run_dir = build_run(cls.ws)
        cls.data = load_run(resolve_run("test 01", workspace=cls.ws))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_run_is_found_by_any_spelling_and_status_is_read(self):
        for name in ("test01", "test_01", "Test 01"):
            self.assertEqual(resolve_run(name, workspace=self.ws).path, self.run_dir)
        self.assertEqual(self.data.unsolved, ["F00_s003", "F00_s004"])
        self.assertEqual(self.data.status_of("F00_s004"), "infeasible")
        self.assertEqual(self.data.status_of("F00_s003"), "no_solution")

    def test_inputs_are_listed_by_every_set_and_year(self):
        listing = allsets.list_by_set(self.data, "Input", write_all_csv=False)
        commodity = listing.family("Commodity")
        self.assertIn("SpecifiedAnnualDemand", set(commodity.Parameter))
        self.assertTrue(set(YEARS) <= set(commodity.columns))
        summary = listing.summary.set_index("Parameter")
        # read from each future's own CSV, so the FUEL / EMISSION keyed parameters are there
        self.assertIn("InputActivityRatio", summary.index)
        self.assertIn("Emission", summary.loc["AnnualEmissionLimit", "Set_Families"])
        self.assertEqual(summary.loc["InputActivityRatio", "Rows_That_Differ"], 0)

    def test_infeasibility_blames_the_lever_that_breaks_a_bound(self):
        p = perturbed.analyse(self.data)
        listing = allsets.list_by_set(self.data, "Input", write_all_csv=False)
        dx = infeasibility.diagnose(self.data, p, listing)
        verdict = dx.summary.set_index("Future_ID")
        self.assertIn("Likely", verdict.loc["F00_s004", "Verdict"])
        self.assertEqual(verdict.loc["F00_s004", "Suspect_Levers"], "X2")
        self.assertNotIn("Likely", verdict.loc["F00_s003", "Verdict"])
        conflicts = dx.conflicts[dx.conflicts.Future_ID == "F00_s004"]
        self.assertTrue(conflicts.New_In_This_Future.all())
        x3 = dx.levers[(dx.levers.Future_ID == "F00_s004") & (dx.levers.Lever == "X3")].iloc[0]
        self.assertEqual(x3.Score, 0)                           # cost lever, never blamed

    def test_script_writes_everything_into_the_run_diagnosis_folder(self):
        run_diagnosis.main(["--workspace", str(self.ws), "--run", "test01", "--no-open"])
        out = self.run_dir / "diagnosis"
        for name in ("Unsolved_Futures.csv", "Perturbed_Inputs.xlsx", "Inputs_by_Set.xlsx",
                     "Inputs_by_Set_all_years.csv", "Outputs_by_Set.xlsx",
                     "Infeasibility_Diagnosis.xlsx", "Perturbed_Inputs_Dashboard.html"):
            self.assertTrue((out / name).is_file(), name)
        unsolved = pd.read_csv(out / "Unsolved_Futures.csv")
        self.assertEqual(list(unsolved.Future_ID), ["F00_s003", "F00_s004"])


if __name__ == "__main__":
    unittest.main()
