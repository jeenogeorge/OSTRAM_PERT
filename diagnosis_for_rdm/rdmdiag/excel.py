"""Write several tables into one formatted workbook."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F4E46")
HEAD_FONT = Font(color="FFFFFF", bold=True)
RED = PatternFill("solid", fgColor="F6D3CF")
AMBER = PatternFill("solid", fgColor="FBE7C6")
GREEN = PatternFill("solid", fgColor="DCEFE3")
GREY = PatternFill("solid", fgColor="EEEEEE")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILLS = {"red": RED, "amber": AMBER, "green": GREEN, "grey": GREY}


@dataclass
class Sheet:
    name: str
    table: pd.DataFrame
    intro: list[str] = field(default_factory=list)   # lines above the table
    row_fill: object = None      # callable(row: dict) -> "red" / "amber" / "green" / "grey" / None
    freeze_cols: int = 1
    max_width: int = 48


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, (np.floating, float)):
        return None if not np.isfinite(v) else round(float(v), 6)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if v is pd.NA or (isinstance(v, float) and np.isnan(v)):
        return None
    return v


def write_workbook(path: Path, sheets: list[Sheet]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh.name[:31])
        r = 1
        for line in sh.intro:
            ws.cell(r, 1, line)
            if r == 1:
                ws.cell(r, 1).font = Font(bold=True, size=13)
            r += 1
        if sh.intro:
            r += 1
        table = sh.table if sh.table is not None else pd.DataFrame()
        if table.empty and not len(table.columns):
            ws.cell(r, 1, "(nothing to show)")
            continue
        head_row = r
        for j, col in enumerate(table.columns, 1):
            c = ws.cell(r, j, str(col))
            c.fill, c.font, c.border = HEAD, HEAD_FONT, BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        widths = [len(str(c)) for c in table.columns]
        for rec in table.to_dict("records"):
            r += 1
            fill = FILLS.get(sh.row_fill(rec)) if sh.row_fill else None
            for j, col in enumerate(table.columns, 1):
                v = _cell_value(rec[col])
                c = ws.cell(r, j, v)
                c.border = BORDER
                if fill is not None:
                    c.fill = fill
                widths[j - 1] = max(widths[j - 1], len(str(v)) if v is not None else 0)
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(8, w + 2), sh.max_width)
        ws.freeze_panes = ws.cell(head_row + 1, sh.freeze_cols + 1)
        if len(table):
            ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(table.columns))}{r}"
    wb.save(path)
    return path
